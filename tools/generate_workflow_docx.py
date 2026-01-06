from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt


def add_paragraph(doc: Document, text: str, *, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_doc() -> Document:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    title = doc.add_paragraph()
    r = title.add_run("VIET-EXPRESS – LUỒNG HOẠT ĐỘNG XỬ LÝ ĐƠN HÀNG\n(Request → Parcel → Trip → Payment)\nVÀ VAI TRÒ THEO ROLE")
    r.bold = True
    r.font.size = Pt(16)

    add_paragraph(doc, f"Ngày tạo tài liệu: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    add_paragraph(doc, "Nguồn mô tả: tổng hợp từ codebase Spring Boot/Thymeleaf trong workspace (Controllers/Services/Entities/SecurityConfig).", italic=True)

    doc.add_paragraph()

    doc.add_heading("1. Mục tiêu & phạm vi", level=1)
    add_bullets(
        doc,
        [
            "Mục tiêu: mô tả chi tiết luồng xử lý khi phát sinh request đơn hàng từ khách hàng, từ lúc tạo đơn đến khi giao hàng và thanh toán.",
            "Phạm vi: các module Request, Parcel (kiện hàng), Trip (chuyến xe), Payment (thanh toán) và Tracking (ParcelAction).",
            "Lưu ý thuật ngữ: Shipper trong hệ thống đồng thời là Driver (tài xế).",
        ],
    )

    doc.add_heading("2. Các role trong hệ thống (tổng quan)", level=1)

    doc.add_heading("2.1 CUSTOMER (người dùng khách hàng: Sender/Receiver)", level=2)
    add_bullets(
        doc,
        [
            "Tạo đơn hàng (Request) với địa chỉ người gửi/người nhận, mô tả hàng hóa, dịch vụ, COD.",
            "Receiver xác nhận hoặc từ chối đơn (chỉ khi đơn đang PENDING).",
            "Theo dõi trạng thái đơn/kiện, xem lịch sử tracking (ParcelAction).",
            "Xem các khoản thanh toán (Payment) mà mình có trách nhiệm trả (lọc theo payerType).",
        ],
    )
    add_paragraph(doc, "Đường dẫn UI tiêu biểu:")
    add_bullets(
        doc,
        [
            "/customer/create-order (tạo đơn)",
            "/customer/orders/{id} (chi tiết đơn, xác nhận/từ chối nếu là receiver)",
            "/customer/orders/{id}/payments (danh sách payment theo đơn)",
            "/tracking hoặc /public/** (tracking công khai theo requestCode – tuỳ cấu hình UI)",
        ],
    )

    doc.add_heading("2.2 ADMIN (điều phối vận hành cấp cao)", level=2)
    add_bullets(
        doc,
        [
            "Xem danh sách các request trong hệ thống.",
            "Gán (assign) Manager phụ trách một request.",
        ],
    )
    add_paragraph(doc, "Đường dẫn UI tiêu biểu:")
    add_bullets(doc, ["/admin/request", "/admin/request/{id}", "/admin/request/{id}/assign-manager"])

    doc.add_heading("2.3 MANAGER (quản lý đơn, điều phối kho/chuyến/thu phí)", level=2)
    add_bullets(
        doc,
        [
            "Xem các request được gán cho mình.",
            "Chốt đơn (CONFIRMED) sau khi receiver xác nhận (RECEIVER_CONFIRMED).",
            "Giao việc cho Staff xử lý request (tạo parcels/hàng hóa).",
            "Tạo Trip (PICKUP/TRANSFER/DELIVERY/RETURN), gán shipper/driver và xếp parcels lên chuyến.",
            "Tạo Payment, cập nhật paidAmount và status (có lịch sử PaymentTransaction).",
        ],
    )
    add_paragraph(doc, "Đường dẫn UI tiêu biểu:")
    add_bullets(
        doc,
        [
            "/manager/requests (danh sách request được gán)",
            "/manager/requests/{id} (chi tiết request)",
            "/manager/requests/{id}/confirm (chốt đơn)",
            "/manager/requests/{id}/assign-staff (giao staff)",
            "/manager/requests/{id}/trips (quản lý trips của đơn)",
            "/manager/trip-planning (xếp parcels lên trip)",
            "/manager/requests/{id}/payments & /manager/payments (quản lý payment)",
        ],
    )

    doc.add_heading("2.4 STAFF (nhân viên kho/nhân viên xử lý đơn)", level=2)
    add_bullets(
        doc,
        [
            "Xem các request được giao (assignedStaff).",
            "Tạo 1 hoặc nhiều Parcel cho một request (tạo theo danh sách hoặc bulk).",
            "Thao tác kho: nhập kho (check-in), xuất kho (check-out).",
            "Theo dõi danh sách kiện trong kho đang IN_WAREHOUSE.",
        ],
    )
    add_paragraph(doc, "Đường dẫn UI tiêu biểu:")
    add_bullets(
        doc,
        [
            "/staff/requests (request được giao)",
            "/staff/requests/{id} (tạo parcel)",
            "/staff/parcels (quản lý parcel; checkin/checkout)",
            "/staff/warehouse (xem kiện trong kho hiện tại)",
        ],
    )

    doc.add_heading("2.5 SHIPPER / DRIVER (tài xế)", level=2)
    add_bullets(
        doc,
        [
            "Xem danh sách trips được gán.",
            "Bắt đầu chuyến (start) và hoàn thành chuyến (complete).",
            "Cập nhật trạng thái kiện trong chuyến (nếu cần).",
        ],
    )
    add_paragraph(doc, "Đường dẫn UI tiêu biểu:")
    add_bullets(doc, ["/shipper/trips", "/shipper/trip/{id}", "/shipper/trip/{id}/start", "/shipper/trip/{id}/complete"])

    doc.add_heading("3. Dữ liệu chính & trạng thái (theo Entities)", level=1)

    doc.add_heading("3.1 CustomerRequest (đơn hàng / request)", level=2)
    add_paragraph(doc, "Các trạng thái (enum RequestStatus):")
    add_bullets(
        doc,
        [
            "PENDING: Customer vừa tạo đơn, chờ receiver xác nhận.",
            "RECEIVER_CONFIRMED: Receiver đã xác nhận, chờ Manager chốt đơn.",
            "CONFIRMED: Manager đã chốt đơn (thường sau RECEIVER_CONFIRMED).",
            "PICKUP_ASSIGNED, PICKED_UP, IN_WAREHOUSE, IN_TRANSIT, OUT_FOR_DELIVERY, DELIVERED: các trạng thái mở rộng cho các bước vận hành.",
            "COMPLETED: hoàn tất toàn bộ quy trình (giao + thanh toán) – hiện trạng codebase chưa thấy auto-set, có thể cập nhật thủ công.",
            "FAILED, RETURNED, CANCELLED: lỗi/hoàn trả/huỷ.",
        ],
    )

    doc.add_heading("3.2 Parcel (kiện hàng)", level=2)
    add_paragraph(doc, "Các trạng thái (enum ParcelStatus):")
    add_bullets(
        doc,
        [
            "CREATED: kiện mới tạo.",
            "PICKED_UP: đã lấy hàng.",
            "IN_WAREHOUSE: đang nằm trong kho.",
            "IN_TRANSIT: đang vận chuyển.",
            "OUT_FOR_DELIVERY: đang đi giao.",
            "DELIVERED: đã giao thành công.",
            "FAILED: giao thất bại.",
            "RETURNED: hoàn hàng.",
        ],
    )
    add_paragraph(doc, "Trường quan trọng:")
    add_bullets(doc, ["currentLocation (địa điểm hiện tại)", "currentTrip (chuyến đang gán)", "currentShipper (tài xế đang giữ)"])

    doc.add_heading("3.3 Trip (chuyến xe)", level=2)
    add_paragraph(doc, "Các loại chuyến (TripType): PICKUP, TRANSFER, DELIVERY, RETURN")
    add_paragraph(doc, "Trạng thái (TripStatus): CREATED, ASSIGNED, IN_PROGRESS, COMPLETED, CANCELLED")
    add_paragraph(doc, "Ghi chú quan trọng từ code:")
    add_bullets(
        doc,
        [
            "Manager có màn hình trip-planning để ‘xếp hàng lên chuyến’: parcels phù hợp thường là IN_WAREHOUSE tại startLocation.",
            "Khi trip chuyển sang IN_PROGRESS (bằng cập nhật trạng thái trip ở Manager), hệ thống tự động tạo COD payment theo trip cho từng request (nếu có COD).",
            "Khi shipper start trip, hệ thống cũng set trip=IN_PROGRESS và set parcels=IN_TRANSIT.",
        ],
    )

    doc.add_heading("3.4 Payment (thanh toán)", level=2)
    add_paragraph(doc, "PaymentType: SHIPPING_FEE, COD, DEPOSIT")
    add_paragraph(doc, "PaymentScope: FULL_REQUEST (toàn đơn), PER_TRIP (theo chuyến)")
    add_paragraph(doc, "PaymentStatus:")
    add_bullets(
        doc,
        [
            "UNPAID, PARTIALLY_PAID, PAID: nhóm trạng thái thanh toán phí.",
            "COLLECTED_FROM_RECEIVER: đã thu COD từ người nhận.",
            "PAID_TO_SENDER: đã trả COD cho người gửi.",
            "REFUNDED: đã hoàn tiền.",
        ],
    )
    add_paragraph(doc, "Lịch sử thay đổi payment:")
    add_bullets(
        doc,
        [
            "PaymentTransaction lưu lịch sử (STATUS_CHANGE, IN, …).",
            "Manager cập nhật payment thường gọi changePaymentStatus(...) để vừa cập nhật status vừa ghi transaction.",
            "Customer chỉ xem các payment phù hợp payerType (Sender/Receiver).",
        ],
    )

    doc.add_heading("3.5 Tracking (ParcelAction / ActionType)", level=2)
    add_bullets(
        doc,
        [
            "Tracking được ghi dưới dạng ParcelAction (gắn với request hoặc parcel).",
            "Các actionCode thường gặp: CREATED, RECEIVER_CONFIRMED, CONFIRMED, IN_WAREHOUSE, IN_TRANSIT, DELIVERED, FAILED, RETURNED, PICKED_UP, LOCATION_CHANGE…",
            "TrackingService.logAction(...) cho phép lưu fromLocation/toLocation, actorUser và note.",
        ],
    )

    doc.add_heading("4. Luồng xử lý chuẩn: từ tạo đơn đến giao hàng", level=1)

    doc.add_heading("Bước 1 – Customer tạo request (tạo đơn)", level=2)
    add_bullets(
        doc,
        [
            "Customer nhập thông tin gửi/nhận, mô tả hàng hóa, dịch vụ, COD.",
            "Hệ thống tạo Location cho sender (SENDER) và receiver (RECEIVER).",
            "Tạo CustomerRequest với status=PENDING.",
            "Ghi tracking action CREATED (gắn với request).",
        ],
    )

    doc.add_heading("Bước 2 – Receiver xác nhận (hoặc từ chối)", level=2)
    add_bullets(
        doc,
        [
            "Receiver chỉ xác nhận khi đơn đang PENDING.",
            "Nếu xác nhận: status → RECEIVER_CONFIRMED + tracking action RECEIVER_CONFIRMED.",
            "Nếu từ chối: status → CANCELLED.",
        ],
    )

    doc.add_heading("Bước 3 – Admin gán Manager phụ trách request", level=2)
    add_bullets(
        doc,
        [
            "Admin chọn manager và gán vào trường assignedManager + managerAssignedAt.",
            "Từ thời điểm này, Manager chỉ thấy đơn được gán cho mình trong màn hình /manager/requests.",
        ],
    )

    doc.add_heading("Bước 4 – Manager chốt đơn và giao Staff", level=2)
    add_bullets(
        doc,
        [
            "Manager chốt đơn (confirm) chỉ khi trạng thái đang RECEIVER_CONFIRMED.",
            "Khi chốt: status → CONFIRMED + tracking action CONFIRMED.",
            "Manager gán staff xử lý đơn (assignedStaff + assignedAt) để staff tạo parcels.",
            "Nếu tuyến đường (Route) chưa có thì shippingFee/estimatedDeliveryTime có thể null (chờ manager cấu hình tuyến/route).",
        ],
    )

    doc.add_heading("Bước 5 – Staff tạo parcels (kiện hàng) từ request", level=2)
    add_bullets(
        doc,
        [
            "Staff vào request được giao và tạo 1-n parcels (từng món hàng/kiện).",
            "Parcel status ban đầu: CREATED.",
            "currentLocation khi tạo: ưu tiên location kho của staff (nếu staff có kho), nếu không có thì fallback về senderLocation.",
            "Ghi tracking action CREATED cho từng parcel.",
        ],
    )

    doc.add_heading("Bước 6 – Vận hành kho: nhập kho / xuất kho", level=2)
    add_bullets(
        doc,
        [
            "Nhập kho (check-in): Parcel status → IN_WAREHOUSE; currentLocation → kho của staff; reset currentTrip/currentShipper; ghi tracking IN_WAREHOUSE.",
            "Xuất kho (check-out): Parcel status → IN_TRANSIT; ghi tracking IN_TRANSIT.",
        ],
    )

    doc.add_heading("Bước 7 – Manager tạo Trip, xếp hàng và gán tài xế (shipper/driver)", level=2)
    add_bullets(
        doc,
        [
            "Manager tạo Trip với tripType, startLocation, endLocation (và gán shipper nếu có).",
            "Ở màn hình trip-planning, manager xếp parcels (IN_WAREHOUSE tại startLocation) lên trip.",
            "Khi xếp lên trip: parcel.currentTrip=trip và thường set parcel.status=IN_TRANSIT.",
            "Manager có thể cập nhật trạng thái trip. Khi trip chuyển IN_PROGRESS (ở Manager), hệ thống auto tạo COD payment theo trip (nếu parcels có COD).",
        ],
    )

    doc.add_heading("Bước 8 – Shipper/Driver thực thi chuyến", level=2)
    add_bullets(
        doc,
        [
            "Start trip: Trip status → IN_PROGRESS, set startedAt; cập nhật tất cả parcels trong trip → IN_TRANSIT.",
            "Complete trip: Trip status → COMPLETED, set endedAt.",
            "Khi complete, trạng thái parcels phụ thuộc tripType:",
            "- DELIVERY: parcels → DELIVERED",
            "- PICKUP: parcels → PICKED_UP",
            "- TRANSFER/RETURN: parcels → IN_WAREHOUSE (về điểm đến để staff nhập kho/tiếp tục xử lý)",
        ],
    )

    doc.add_heading("Bước 9 – Thanh toán & hoàn tất", level=2)
    add_bullets(
        doc,
        [
            "Manager có thể tạo các khoản Payment (SHIPPING_FEE/COD/DEPOSIT) theo FULL_REQUEST hoặc PER_TRIP, với payerType/receiverType phù hợp.",
            "Manager cập nhật paidAmount và status; lịch sử status được ghi vào PaymentTransaction.",
            "Các trạng thái COD (COLLECTED_FROM_RECEIVER/PAID_TO_SENDER) đã được định nghĩa; hiện trạng codebase chủ yếu hỗ trợ cập nhật thủ công qua màn hình manager.",
            "RequestStatus.COMPLETED tồn tại trong enum nhưng hiện chưa thấy auto-set trong code; thường được manager cập nhật khi đã giao + đã thu/chi đủ.",
        ],
    )

    doc.add_heading("5. Ví dụ chi tiết (multi-kho, multi-trip)", level=1)

    add_paragraph(doc, "Giả định địa điểm:")
    add_bullets(
        doc,
        [
            "SenderLocation: TP.HCM (SENDER)",
            "Kho HCM: WAREHOUSE_HCM", 
            "Kho Đà Nẵng: WAREHOUSE_DN",
            "Kho Hà Nội: WAREHOUSE_HN",
            "ReceiverLocation: Hà Nội (RECEIVER)",
        ],
    )

    add_paragraph(doc, "Giả định đơn có 2 kiện: Parcel A (COD 500k) và Parcel B (COD 0).")

    doc.add_heading("5.1 Dòng thời gian (ai làm gì, hệ thống cập nhật gì)", level=2)
    add_bullets(
        doc,
        [
            "(Customer/Sender) tạo đơn: Request= PENDING, tracking CREATED.",
            "(Customer/Receiver) xác nhận: Request → RECEIVER_CONFIRMED, tracking RECEIVER_CONFIRMED.",
            "(Admin) gán Manager M1: request.assignedManager=M1.",
            "(Manager M1) chốt đơn: Request → CONFIRMED, tracking CONFIRMED.",
            "(Manager M1) gán Staff S1 (kho HCM).",
            "(Staff S1) tạo Parcel A/B: Parcel status=CREATED; currentLocation thường = kho HCM (nếu staff có kho); tracking CREATED cho từng parcel.",
            "(Staff S1) nhập kho nếu cần: Parcel → IN_WAREHOUSE, currentLocation=kho HCM; tracking IN_WAREHOUSE.",
            "(Manager M1) tạo Trip T1 (TRANSFER) từ kho HCM → kho Đà Nẵng; xếp Parcel A/B lên T1.",
            "(Shipper/Driver D1) start T1: Trip=IN_PROGRESS; parcels=IN_TRANSIT.",
            "(Shipper/Driver D1) complete T1: Trip=COMPLETED; parcels → IN_WAREHOUSE (đến kho Đà Nẵng để tiếp tục trung chuyển).",
            "(Staff kho Đà Nẵng) check-in: parcels IN_WAREHOUSE tại kho Đà Nẵng.",
            "(Manager M1) tạo Trip T2 (TRANSFER) kho Đà Nẵng → kho Hà Nội; xếp Parcel A/B.",
            "(Driver D2) start/complete T2: parcels về IN_WAREHOUSE tại kho Hà Nội.",
            "(Manager M1) tạo Trip T3 (DELIVERY) kho Hà Nội → ReceiverLocation; xếp Parcel A/B.",
            "(Driver D3) start/complete T3: parcels → DELIVERED.",
            "(Manager M1) tạo/cập nhật Payment: ",
            "- COD: có thể được auto tạo khi trip chuyển IN_PROGRESS (theo màn hình Manager) hoặc tạo thủ công; cập nhật thu/chi COD theo nghiệp vụ.",
            "- SHIPPING_FEE: tạo payment và thu từ payerType phù hợp.",
            "(Manager M1) khi đã giao + thu phí xong: cập nhật Request → COMPLETED (nếu quy trình vận hành yêu cầu).",
        ],
    )

    doc.add_heading("5.2 Checkpoints kiểm soát (điểm dễ sai)", level=2)
    add_bullets(
        doc,
        [
            "Receiver chỉ xác nhận khi PENDING; nếu không xác nhận thì Manager không chốt được (trừ force-confirm).",
            "Trip-planning chỉ xếp được parcels IN_WAREHOUSE tại startLocation: vì vậy cần quy trình kho hợp lý trước khi tạo trip.",
            "Khi nhập kho, hệ thống reset currentTrip/currentShipper để tránh lẫn lộn chuyến cũ.",
            "Payment: không nên tạo tổng expectedAmount vượt quá (shippingFee + codAmount) còn lại; ManagerPaymentController có validation cho việc này.",
        ],
    )

    doc.add_heading("6. Mapping nhanh: ai thao tác ở đâu (tóm tắt)", level=1)
    add_bullets(
        doc,
        [
            "Customer: tạo đơn, receiver xác nhận, theo dõi & xem payment theo vai trò sender/receiver.",
            "Admin: gán manager phụ trách request.",
            "Manager: chốt đơn, gán staff, tạo trip, xếp hàng, gán driver, tạo & cập nhật payment.",
            "Staff: tạo parcels, nhập/xuất kho, chuẩn bị hàng cho các trip.",
            "Shipper/Driver: chạy trip (start/complete), cập nhật trạng thái parcels trong trip.",
        ],
    )

    doc.add_paragraph()
    add_paragraph(doc, "Kết thúc tài liệu.")

    return doc


def main() -> None:
    workspace = Path(__file__).resolve().parents[1]
    out_dir = workspace / "docs"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "luong-hoat-dong-role-workflow.docx"

    doc = build_doc()
    doc.save(out_path)
    print(f"Wrote: {out_path}")


if __name__ == "__main__":
    main()
